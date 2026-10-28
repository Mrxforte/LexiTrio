"""Extract plain text from uploaded dictionary documents."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from pathlib import Path

from django.conf import settings

SUPPORTED_EXTENSIONS = (".txt", ".text", ".md", ".markdown", ".csv", ".tsv", ".pdf", ".docx")

_CONTROL_RE = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_BLANK_LINES_RE = re.compile(r"\n{3,}")
_TRAILING_SPACE_RE = re.compile(r"[ \t]+\n")


class DocumentError(Exception):
    """The uploaded document cannot be converted to text."""


@dataclass(frozen=True)
class ExtractedDocument:
    text: str
    file_name: str
    extension: str
    size_bytes: int
    page_count: int | None = None

    @property
    def characters(self) -> int:
        return len(self.text)


def normalize_text(value: str) -> str:
    if not value:
        return ""
    text = value.replace("\r\n", "\n").replace("\r", "\n")
    text = _CONTROL_RE.sub("", text)
    text = _TRAILING_SPACE_RE.sub("\n", text)
    text = _BLANK_LINES_RE.sub("\n\n", text)
    return text.strip()


def extract_document(upload) -> ExtractedDocument:
    file_name = Path(getattr(upload, "name", "") or "document").name
    extension = Path(file_name).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentError(
            "Qo'llab-quvvatlanmaydigan fayl turi. Ruxsat etilgan: "
            + ", ".join(SUPPORTED_EXTENSIONS)
        )

    max_bytes = int(getattr(settings, "TRANSLATION_MAX_UPLOAD_BYTES", 10 * 1024 * 1024))
    size_bytes = int(getattr(upload, "size", 0) or 0)
    if size_bytes > max_bytes:
        raise DocumentError(
            f"Fayl hajmi juda katta ({size_bytes} bayt). Ruxsat etilgan limit: {max_bytes} bayt."
        )

    raw = upload.read()
    page_count: int | None = None

    if extension == ".pdf":
        text, page_count = _extract_pdf(raw)
    elif extension == ".docx":
        text = _extract_docx(raw)
    elif extension in {".csv", ".tsv"}:
        text = _extract_delimited(raw, extension)
    else:
        text = _decode_bytes(raw)

    text = normalize_text(text)
    if not text:
        raise DocumentError("Fayldan matn topilmadi.")

    return ExtractedDocument(
        text=text,
        file_name=file_name,
        extension=extension,
        size_bytes=size_bytes or len(raw),
        page_count=page_count,
    )


def _decode_bytes(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1251", "koi8-r"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


def _extract_pdf(data: bytes) -> tuple[str, int]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise DocumentError(
            "PDF o'qish uchun 'pypdf' kutubxonasi kerak: pip install -r requirements.txt"
        ) from exc

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception as exc:
        raise DocumentError("PDF faylni o'qib bo'lmadi (buzilgan yoki shifrlangan).") from exc

    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")
        except Exception as exc:
            raise DocumentError("PDF fayl parol bilan himoyalangan.") from exc

    pages: list[str] = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    return "\n\n".join(pages), len(reader.pages)


def _extract_docx(data: bytes) -> str:
    try:
        from docx import Document as DocxDocument
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise DocumentError(
            "DOCX o'qish uchun 'python-docx' kutubxonasi kerak: pip install -r requirements.txt"
        ) from exc

    try:
        document = DocxDocument(io.BytesIO(data))
    except Exception as exc:
        raise DocumentError("DOCX faylni o'qib bo'lmadi (buzilgan fayl).") from exc

    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cell for cell in cells if cell))
    return "\n".join(parts)


def _extract_delimited(data: bytes, extension: str) -> str:
    text = _decode_bytes(data)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = None

    reader = (
        csv.reader(io.StringIO(text), dialect)
        if dialect
        else csv.reader(io.StringIO(text), delimiter="\t" if extension == ".tsv" else ",")
    )

    lines = []
    for row in reader:
        cells = [cell.strip() for cell in row]
        if any(cells):
            lines.append(" | ".join(cell for cell in cells if cell))
    return "\n".join(lines)
